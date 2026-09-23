"""Browser regression on synthetic data in a temporary DB; no model inference or user data."""

import argparse
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request
import wave
from pathlib import Path

from playwright.sync_api import expect, sync_playwright

from autoprotocol import analysis, jobs, review
from autoprotocol.pipeline.extract import validate_output
from autoprotocol.storage import connect, initialize


def seed(folder):
    database = folder / "autoprotocol.sqlite3"
    initialize(database)
    root = Path(__file__).resolve().parents[1]
    snapshot = json.loads((root / "tests/fixtures/extraction_ru.json").read_text(encoding="utf-8"))
    for segment in snapshot["segments"]:
        segment.update(review_flags=[], words=[])
    snapshot["segments"][1]["text"] += " Выпуск составил 94% от плана. Поставщик задержал сырьё."
    result = folder / "result.json"
    result.write_text(json.dumps(snapshot, ensure_ascii=False), encoding="utf-8")
    audio = folder / "audio.wav"
    with wave.open(str(audio), "wb") as stream:
        stream.setparams((1, 2, 16000, 0, "NONE", "NONE"))
        stream.writeframes(b"\0\0" * 16000 * 7)
    jobs.enqueue(
        database,
        {
            "id": "browser-check",
            "request_key": "browser-check",
            "title": "Синтетический тест редактора",
            "occurred_on": snapshot["occurred_on"],
            "timezone": snapshot["timezone"],
            "now": time.time(),
            "duration_seconds": 7,
            "source_path": str(audio),
            "audio_sha256": "test",
        },
    )
    with connect(database) as db:
        db.execute(
            "UPDATE meetings SET status='ready',stage='transcript_ready',"
            "result_path=?,audio_path=?",
            (str(result), str(audio)),
        )
    current = review.get_result(database, "browser-check")
    analysis.enqueue(
        database,
        "browser-check",
        analysis.AnalysisRequest(
            expected_revision=0, source_digest=current["review"]["source_digest"]
        ),
    )
    evidence = [
        {"segment_id": "seg_1", "quote": snapshot["segments"][0]["text"], "role": "initial"}
    ]
    fixture = {
        "actions": [
            {
                "task": "Подготовить отчёт",
                "assignee_text": "Ерлан",
                "due_text": "завтра",
                "due_kind": "date",
                "status": "agreed",
                "issuer_segment_id": "seg_1",
                "evidence": evidence,
            }
        ],
        "summary": [{"kind": "action", "text": "Подготовить отчёт", "evidence": evidence}],
        "reports": [
            {
                "direction": "Производство",
                "indicator": "94% от плана",
                "problem": "Поставщик задержал сырьё",
                "evidence": [
                    {
                        "segment_id": "seg_2",
                        "quote": snapshot["segments"][1]["text"],
                        "role": "initial",
                    }
                ],
            }
        ],
    }
    job = analysis.claim(database)
    assert analysis.finish(database, job, result=validate_output(json.dumps(fixture), snapshot))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--channel", default="chrome", choices=["chrome", "msedge", "chromium"])
    args = parser.parse_args()
    artifacts = Path("data/browser-review-check").resolve()
    artifacts.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="review-browser-", dir=artifacts) as directory:
        folder = Path(directory)
        seed(folder)
        with socket.socket() as sock:
            sock.bind(("127.0.0.1", 0))
            port = sock.getsockname()[1]
        base = f"http://127.0.0.1:{port}"
        env = {**os.environ, "AUTOPROTOCOL_DATA_DIR": str(folder)}
        process = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "autoprotocol.main:create_app",
                "--factory",
                "--host",
                "127.0.0.1",
                "--port",
                str(port),
                "--no-access-log",
            ],
            env=env,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        try:
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
            deadline = time.monotonic() + 15
            while True:
                try:
                    with opener.open(base + "/health/live", timeout=1):
                        break
                except OSError:
                    if time.monotonic() >= deadline:
                        raise RuntimeError("Test server did not start") from None
                    time.sleep(0.1)
            with sync_playwright() as playwright:
                browser = playwright.chromium.launch(channel=args.channel, headless=True)
                context = browser.new_context(viewport={"width": 1280, "height": 900})
                errors = []
                context.on(
                    "page",
                    lambda page: page.on("pageerror", lambda error: errors.append(str(error))),
                )
                page = context.new_page()
                page.goto(base + "/#browser-check")
                action = page.locator('[data-action-id="action_0001"]')
                expect(action).to_be_visible()
                action.locator('[data-field="task"]').fill("Проверенный отчёт — Қазақстан")
                action.locator('[data-field="assignee_text"]').fill("Әлия")
                action.locator('[data-field="due_text"]').fill("2026-10-15")
                action.locator(".item-reviewed").check()
                expect(page.locator("#review-fields")).to_have_js_property("disabled", True)
                page.evaluate("refreshAnalysis(selected)")
                expect(action.locator('[data-field="task"]')).to_have_value(
                    "Проверенный отчёт — Қазақстан"
                )
                page.locator("#save-analysis").click()
                expect(page.locator("#analysis-review-message")).to_contain_text("Версия 1")
                page.reload()
                expect(action.locator('[data-field="assignee_text"]')).to_have_value("Әлия")
                expect(action.locator(".item-reviewed")).to_be_checked()
                expect(action).to_contain_text("2026-10-15")
                action.get_by_text("Исходный ответ модели", exact=True).click()
                expect(action).to_contain_text("Ерлан")
                action.get_by_text("Источники в записи", exact=True).click()
                action.locator("details button").click()
                expect(page.locator("#meeting-audio")).to_have_js_property("paused", False)
                page.locator("#analysis-history summary").click()
                expect(page.locator("#analysis-history-list")).to_contain_text("Әлия")
                page.screenshot(path=str(artifacts / "desktop.png"), full_page=True)
                page.locator("#analysis-heading").scroll_into_view_if_needed()
                page.screenshot(path=str(artifacts / "desktop-viewport.png"))

                second = context.new_page()
                second.goto(base + "/#browser-check")
                second_action = second.locator('[data-action-id="action_0001"]')
                expect(second_action).to_be_visible()
                action.locator('[data-field="task"]').fill("Черновик в первой вкладке")
                expect(action.locator(".item-reviewed")).not_to_be_checked()
                second_action.locator('[data-field="task"]').fill("Правка во второй вкладке")
                second.locator("#save-analysis").click()
                expect(second.locator("#analysis-review-message")).to_contain_text("Версия 2")
                page.evaluate("refreshAnalysis(selected)")
                expect(action.locator('[data-field="task"]')).to_have_value(
                    "Черновик в первой вкладке"
                )
                page.locator("#save-analysis").click()
                expect(page.locator("#analysis-review-message")).to_contain_text("другой вкладке")
                page.once("dialog", lambda dialog: dialog.accept())
                page.locator("#reload-analysis").click()
                expect(action.locator('[data-field="task"]')).to_have_value(
                    "Правка во второй вкладке"
                )
                summary = page.locator("[data-summary-id]")
                summary.locator(".item-excluded").check()
                page.locator("#save-analysis").click()
                expect(page.locator("#analysis-review-message")).to_contain_text("Версия 3")
                page.reload()
                expect(summary.locator(".item-excluded")).to_be_checked()
                report = page.locator("[data-report-id]").first
                report.locator('[data-field="indicator"]').fill("94% от плана — проверено")
                report.locator(".item-reviewed").check()
                expect(page.locator("#download-docx")).to_be_disabled()
                page.locator("#save-analysis").click()
                expect(page.locator("#analysis-review-message")).to_contain_text("Версия 4")
                page.reload()
                expect(report.locator('[data-field="indicator"]')).to_have_value(
                    "94% от плана — проверено"
                )
                for template in ("1", "2"):
                    page.locator("#export-template").select_option(template)
                    page.locator("#export-organization").fill("Синтетическая организация")
                    with page.expect_download() as downloaded:
                        page.locator("#download-docx").click()
                    download = downloaded.value
                    path = artifacts / f"protocol-{template}.docx"
                    download.save_as(path)
                    from docx import Document

                    document = Document(path)
                    assert any(
                        "Правка во второй вкладке" in cell.text
                        for table in document.tables
                        for row in table.rows
                        for cell in row.cells
                    )
                    assert not any("Подготовить отчёт [" in p.text for p in document.paragraphs)
                    if template == "2":
                        assert "94% от плана — проверено" in document.tables[0].cell(1, 1).text
                page.set_viewport_size({"width": 390, "height": 844})
                page.locator("#analysis-heading").scroll_into_view_if_needed()
                page.screenshot(path=str(artifacts / "mobile.png"), full_page=True)
                page.screenshot(path=str(artifacts / "mobile-viewport.png"))
                overflow = page.evaluate("""[...document.querySelectorAll('body *')]
                    .filter(el => el.getBoundingClientRect().right + scrollX > innerWidth)
                    .map(el => ({tag: el.tagName, id: el.id, class: el.className,
                        width: el.getBoundingClientRect().width,
                        right: el.getBoundingClientRect().right,
                        scrollX, innerWidth, rootWidth: document.documentElement.scrollWidth}))""")
                assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), (
                    overflow
                )

                # Changing the saved transcript makes existing analysis read-only.
                second.reload()
                text = second.locator(".segment-text").first
                expect(text).to_be_visible()
                corrected = text.input_value() + " Исправлено вручную."
                text.fill(corrected)
                second.locator(".segment-reviewed").first.check()
                second.locator(".save-review").first.click()
                expect(second.locator("#review-message")).to_contain_text("Версия 1")
                second.reload()
                expect(text).to_have_value(corrected)
                expect(second.locator(".segment-reviewed").first).to_be_checked()
                page.evaluate("refreshAnalysis(selected)")
                expect(page.locator("#analysis-review-fields")).to_have_js_property(
                    "disabled", True
                )
                expect(page.locator("#analysis-message")).to_contain_text(
                    "Версия анализа отличается"
                )
                assert not errors, errors
                browser.close()
            print(
                "Browser checks passed: save/reload, original/evidence/audio, history, polling, "
                "two-tab conflict, report edits, both DOCX downloads, exclusion, "
                "stale transcript, mobile width; no JS errors."
            )
        finally:
            process.terminate()
            process.wait(timeout=10)


if __name__ == "__main__":
    main()
