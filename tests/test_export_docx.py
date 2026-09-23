import copy
import hashlib
import io
import json
import sqlite3
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from test_analysis import ready as prepared_meeting  # noqa: F401
from test_analysis_review import extracted as extracted  # noqa: F401

from autoprotocol import analysis, analysis_review, export_docx, review
from autoprotocol.main import create_app
from autoprotocol.pipeline.extract import unknown_report_value, validate_output
from autoprotocol.storage import connect, initialize


def request_for(config, identifier, template="1"):
    result = analysis.latest(config.database_path, "meeting")
    return export_docx.ExportRequest(
        template=template,
        organization="Тестовая ұйым",
        analysis_id=identifier,
        source_digest=result["source_digest"],
        transcript_revision=result["revision"],
        analysis_revision=result["result"]["review"]["revision"],
    )


def texts(document):
    return "\n".join(t.text or "" for t in document.element.iter(qn("w:t")))


@pytest.mark.parametrize("template", ["1", "2"])
def test_download_saved_edits_and_exclusions_without_altering_model(extracted, template):
    config, identifier, original, _ = extracted
    result = analysis.latest(config.database_path, "meeting")["result"]
    action = {key: result["actions"][0][key] for key in analysis_review.ActionEdit.model_fields}
    action.update(
        task="Проверить поставку Қазақстан",
        assignee_text="Әлия",
        reviewed=True,
        due_text=None,
        due_kind="unspecified",
    )
    summary = {key: result["summary"][0][key] for key in analysis_review.SummaryEdit.model_fields}
    summary.update(text="Скрытый ошибочный вывод", excluded=True)
    analysis_review.save(
        config.database_path,
        "meeting",
        identifier,
        analysis_review.AnalysisPatch(expected_revision=0, actions=[action], summary=[summary]),
    )
    request = request_for(config, identifier, template)
    with TestClient(create_app(config), base_url="http://localhost") as client:
        response = client.post("/api/meetings/meeting/export.docx", json=request.model_dump())
        assert response.status_code == 200
        assert response.headers["content-type"] == export_docx.MIME
        assert response.headers["cache-control"] == "no-store"
    document = Document(io.BytesIO(response.content))
    text = texts(document)
    assert "Проверить поставку Қазақстан" in text and "Әлия" in text and "Не указан" in text
    assert "Скрытый ошибочный вывод" not in text
    assert "Самрук" not in text and "Асхат" not in text
    assert "Тестовая ұйым" in text
    assert "Требует проверки" not in text  # Only included action was reviewed.
    assert "Версия правок анализа: 1" in text
    assert original["actions"][0]["evidence"][0]["quote"] in text
    assert "00:00" in text
    with connect(config.database_path) as db:
        assert (
            json.loads(
                db.execute("SELECT result_json FROM analyses WHERE id=?", (identifier,)).fetchone()[
                    0
                ]
            )
            == original
        )
    expected = [4680, 2808, 1872] if template == "1" else [4212, 2808, 2340]
    assert [int(col.get(qn("w:w"))) for col in document.tables[-1]._tbl.tblGrid] == expected
    assert document.sections[0].page_width.twips == 12240
    assert document.sections[0].page_height.twips == 15840


def test_stale_and_foreign_versions_origin_and_bad_template(extracted):
    config, identifier, _, _ = extracted
    payload = request_for(config, identifier).model_dump()
    with TestClient(create_app(config), base_url="http://localhost") as client:
        url = "/api/meetings/meeting/export.docx"
        for field in ("analysis_revision", "transcript_revision"):
            assert client.post(url, json={**payload, field: 9}).status_code == 409
        assert client.post(url, json={**payload, "source_digest": "0" * 64}).status_code == 409
        assert client.post(url, json={**payload, "analysis_id": "foreign"}).status_code == 404
        assert client.post(url, json={**payload, "template": "../../file"}).status_code == 422
        assert (
            client.post(
                url, json=payload, headers={"Origin": "https://elsewhere.example"}
            ).status_code
            == 403
        )
        with connect(config.database_path) as db:
            db.execute("UPDATE analyses SET status='processing'")
        assert client.post(url, json=payload).status_code == 409


def test_topics_partition_saved_transcript_and_assign_each_action_once(extracted):
    config, identifier, _, _ = extracted
    request = request_for(config, identifier)
    request.topics = [
        export_docx.Topic(title="Первая тема", start_segment_id="seg_1"),
        export_docx.Topic(title="Вторая тема", start_segment_id="seg_2"),
    ]
    doc = Document(io.BytesIO(export_docx.export(config.database_path, "meeting", request)))
    assert len(doc.tables) == 2
    assert "Подготовить отчёт" in texts(doc)
    assert "Часть 2. Вторая тема" in texts(doc)
    assert "Поручения не выделены" in doc.tables[1].cell(1, 0).text
    for starts in (("seg_2", "seg_1"), ("seg_1", "seg_1"), ("seg_1", "foreign")):
        request.topics = [export_docx.Topic(title="Тема", start_segment_id=s) for s in starts]
        with pytest.raises(review.ReviewError):
            export_docx.export(config.database_path, "meeting", request)


def test_unknown_date_roles_and_unsafe_xml_are_not_invented(extracted):
    config, identifier, _, _ = extracted
    request = request_for(config, identifier)
    request.roles = [export_docx.Role(speaker_id="SPEAKER_00", role="Директор")]
    with pytest.raises(review.ReviewError, match="подтверждённому"):
        export_docx.export(config.database_path, "meeting", request)
    request.roles = []
    data = export_docx.snapshot(config.database_path, "meeting", request)
    data["meeting"]["occurred_on"] = None
    data["transcript"]["segments"][0]["text"] = "Текст с \x00 символом и <xml> & Ә"
    text = texts(Document(io.BytesIO(export_docx.render(data, request))))
    assert "Дата встречи: не указана" in text
    assert "Текст с � символом и <xml> & Ә" in text


def test_package_parts_and_sanitized_templates_preserve_reference_format():
    manifest = json.loads(
        (Path(__file__).parents[1] / "docs/docx-template-manifest.json").read_text()
    )
    for number in ("1", "2"):
        path = export_docx.TEMPLATES / f"protocol-{number}.docx"
        assert hashlib.sha256(path.read_bytes()).hexdigest() == manifest[number]["template_sha256"]
        with ZipFile(path) as package:
            for part, digest in manifest[number]["preserved_parts"].items():
                assert hashlib.sha256(package.read(part)).hexdigest() == digest
            xml = package.read("word/document.xml").decode()
            assert "Асхат" not in xml and "Ерлан" not in xml and "Самрук" not in xml


def test_reports_validate_exact_quotes_and_missing_indicators():
    fixture = json.loads(
        (Path(__file__).parent / "fixtures/report_cases.json").read_text(encoding="utf-8")
    )[0]["meeting"]
    rows = [
        {
            "direction": "Производство",
            "indicator": "94% от плана",
            "problem": "поставщик задержал сырьё",
            "evidence": [
                {"segment_id": "seg_1", "quote": fixture["segments"][0]["text"], "role": "initial"}
            ],
        },
        {
            "direction": "Обучение",
            "indicator": None,
            "problem": "у 12% сотрудников просрочены сертификаты",
            "evidence": [
                {"segment_id": "seg_2", "quote": fixture["segments"][1]["text"], "role": "initial"}
            ],
        },
    ]
    payload = {"actions": [], "summary": [], "reports": rows}
    result = validate_output(json.dumps(payload), fixture)
    assert result["reports"][1]["indicator"] is None
    assert result["reports"][1]["evidence"][0]["start_ms"] == 4000
    rows[1]["indicator"] = "Показатель пока не назван"
    rows[0]["problem"] = "Поставщик задержал сырьё"
    result = validate_output(json.dumps(payload), fixture)
    assert result["reports"][0]["problem"] == "поставщик задержал сырьё"
    assert result["reports"][1]["indicator"] is None
    assert result["reports"][1]["problem"] == "у 12% сотрудников просрочены сертификаты"
    assert not unknown_report_value("0%")
    assert not unknown_report_value("Проблем нет")
    assert not unknown_report_value("показатель не назван, но известно 12%")
    bad = copy.deepcopy(payload)
    bad["reports"][0]["indicator"] = "100%"
    with pytest.raises(ValueError, match="unsupported_report_field"):
        validate_output(json.dumps(bad), fixture)


def test_schema_upgrade_keeps_existing_manual_history_and_allows_new_run(extracted):
    config, identifier, _, request = extracted
    with connect(config.database_path) as db:
        db.execute("UPDATE analyses SET prompt_version='meeting-evidence-v2'")
    assert analysis.latest(config.database_path, "meeting")["upgrade_available"]
    new_id = analysis.enqueue(config.database_path, "meeting", analysis.AnalysisRequest(**request))
    assert new_id != identifier
    initialize(config.database_path)
    with connect(config.database_path) as db:
        assert db.execute("SELECT count(*) FROM analyses").fetchone()[0] == 2
        assert not db.execute("PRAGMA foreign_key_check").fetchall()


def test_migration_from_populated_v5_preserves_revisions(extracted, tmp_path):
    config, identifier, _, _ = extracted
    current = analysis.latest(config.database_path, "meeting")["result"]
    item = {key: current["actions"][0][key] for key in analysis_review.ActionEdit.model_fields}
    item.update(task="Правка до миграции", reviewed=True)
    analysis_review.save(
        config.database_path,
        "meeting",
        identifier,
        analysis_review.AnalysisPatch(expected_revision=0, actions=[item]),
    )
    legacy = tmp_path / "legacy.sqlite3"
    sql = (Path(__file__).parent / "fixtures/schema_v5.sql").read_text(encoding="utf-8")
    with sqlite3.connect(legacy) as destination, connect(config.database_path) as source:
        destination.executescript(sql)
        for table in ("meetings", "jobs", "analyses", "analysis_reviews", "analysis_revisions"):
            columns = [r[1] for r in destination.execute(f"PRAGMA table_info({table})")]
            for row in source.execute(f"SELECT {','.join(columns)} FROM {table}"):
                destination.execute(
                    f"INSERT INTO {table} VALUES ({','.join('?' for _ in columns)})", tuple(row)
                )
    initialize(legacy)
    initialize(legacy)
    restored = analysis.latest(legacy, "meeting")
    assert restored["result"]["actions"][0]["task"] == "Правка до миграции"
    assert restored["result"]["review"]["revision"] == 1
    assert len(analysis_review.history(legacy, "meeting", identifier)) == 1
    with connect(legacy) as db:
        assert not db.execute("PRAGMA foreign_key_check").fetchall()


def test_report_edit_is_saved_exported_and_exclusion_removes_it(extracted):
    config, identifier, _, _ = extracted
    with connect(config.database_path) as db:
        row = db.execute("SELECT result_json FROM analyses WHERE id=?", (identifier,)).fetchone()
        raw = json.loads(row[0])
        raw["reports"] = [
            {
                "direction": "Синтетический доклад",
                "indicator": None,
                "problem": "Уточнить",
                "evidence": raw["actions"][0]["evidence"],
            }
        ]
        db.execute("UPDATE analyses SET result_json=? WHERE id=?", (json.dumps(raw), identifier))
    payload = analysis_review.AnalysisPatch(
        expected_revision=0,
        reports=[
            {
                "id": "reports_1",
                "direction": "Обучение",
                "indicator": "12%",
                "problem": None,
                "reviewed": True,
                "excluded": False,
            }
        ],
    )
    saved = analysis_review.save(config.database_path, "meeting", identifier, payload)
    assert saved["reports"][0]["original"]["indicator"] is None
    document = Document(
        io.BytesIO(
            export_docx.export(
                config.database_path, "meeting", request_for(config, identifier, "2")
            )
        )
    )
    assert document.tables[0].cell(1, 1).text == "12%"
    assert document.tables[0].cell(1, 2).text == "Не указана"
    payload.expected_revision = 1
    payload.reports[0].excluded = True
    analysis_review.save(config.database_path, "meeting", identifier, payload)
    document = Document(
        io.BytesIO(
            export_docx.export(
                config.database_path, "meeting", request_for(config, identifier, "2")
            )
        )
    )
    assert "12%" not in texts(document)


def test_empty_transcript_has_valid_empty_tables_and_legacy_reports_are_blocked(extracted):
    config, identifier, _, _ = extracted
    request = request_for(config, identifier)
    data = export_docx.snapshot(config.database_path, "meeting", request)
    data["transcript"]["segments"] = []
    data["analysis"].update(actions=[], summary=[], reports=[])
    doc = Document(io.BytesIO(export_docx.render(data, request)))
    assert "Речь не обнаружена" in texts(doc)
    assert len(doc.tables) == 1
    with connect(config.database_path) as db:
        raw = json.loads(
            db.execute("SELECT result_json FROM analyses WHERE id=?", (identifier,)).fetchone()[0]
        )
        raw.pop("reports")
        db.execute("UPDATE analyses SET result_json=? WHERE id=?", (json.dumps(raw), identifier))
    request.template = "2"
    with pytest.raises(review.ReviewError, match="новый анализ"):
        export_docx.export(config.database_path, "meeting", request)
